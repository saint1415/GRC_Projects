#!/usr/bin/env python3
"""Download NIST's official HIPAA Security Rule mappings (OLIR) and write them into the Health Care vertical.

Sources (NIST National Online Informative References, final, posted 2024-03-20, developer NIST):
  HIPAA-Security-Rule-to-SP-800-53-Rev-5.1.1 (OLIR 110)
  HIPAA-Security-Rule-to-Cybersecurity-Framework-v1.1 (OLIR 109)
These are the online mappings that NIST SP 800-66 Rev. 2 Appendix D points to.

Outputs (02_industry-rules/health-care/):
  hipaa-nist-official-mapping.csv   one row per HIPAA citation: official SP 800-53 controls and CSF 1.1 subcategories
  hipaa-security-rule-crosswalk.csv adds/refreshes the official_sp800_53r5_controls column (author CSF 2.0 mapping kept)
Usage: python3 tools/refresh_hipaa_official_mapping.py [--cache DIR]
"""
import collections, csv, hashlib, io, pathlib, re, sys, urllib.request

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
V = ROOT / "02_industry-rules" / "health-care"
BASE = "https://csrc.nist.gov/csrc/media/Projects/olir/documents/submissions/"
FILES = {"sp80053": "SP800-53_SecurityRule_Crosswalk_2024.xlsx", "csf11": "CSFv1-1_SecurityRule_Crosswalk_2024.xlsx"}


def fetch(name, cache):
    if cache and (cache / FILES[name]).exists():
        data = (cache / FILES[name]).read_bytes()
    else:
        with urllib.request.urlopen(BASE + FILES[name], timeout=120) as r:
            data = r.read()
    return openpyxl.load_workbook(io.BytesIO(data), read_only=True), hashlib.sha256(data).hexdigest()


def norm_ctrl(c):
    m = re.match(r"([A-Z]{2})-0*(\d+)(?:\(0*(\d+)\))?$", c.strip())
    return f"{m.group(1)}-{int(m.group(2))}" + (f"({int(m.group(3))})" if m.group(3) else "")


def norm_cit(c):
    c = c.strip()
    return "164.310(c)" if c == "164.310(C)" else c  # one uppercase typo in the NIST file


def main():
    cache = pathlib.Path(sys.argv[sys.argv.index("--cache") + 1]) if "--cache" in sys.argv else None
    wb53, h53 = fetch("sp80053", cache)
    wbcsf, hcsf = fetch("csf11", cache)
    m53, mcsf = collections.defaultdict(set), collections.defaultdict(set)
    for ws in wb53.worksheets:
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if r[0] and r[3]:
                m53[norm_cit(str(r[3]))].add(norm_ctrl(str(r[0])))
    for ws in wbcsf.worksheets:
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if r[0] and len(r) > 2 and r[2] and re.match(r"^[A-Z]{2}\.[A-Z]{2}-\d+$", str(r[0]).strip()):
                mcsf[norm_cit(str(r[2]))].add(str(r[0]).strip())

    def key(c):
        return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", c)]

    rows = list(csv.DictReader(open(V / "hipaa-security-rule-crosswalk.csv")))
    fields = [f for f in rows[0].keys() if f != "official_sp800_53r5_controls"]
    fields.insert(fields.index("csf2_subcategories"), "official_sp800_53r5_controls")
    for r in rows:
        r["official_sp800_53r5_controls"] = "; ".join(sorted(m53.get(r["citation"], []), key=key)) or "Not mapped by NIST"
    with open(V / "hipaa-security-rule-crosswalk.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    with open(V / "hipaa-nist-official-mapping.csv", "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["hipaa_citation", "sp800_53_controls", "csf1_1_subcategories", "sp800_53_version", "csf_version", "source", "sha256_sp800_53_file", "sha256_csf_file"])
        for cit in sorted(set(m53) | set(mcsf), key=key):
            w.writerow([cit, "; ".join(sorted(m53.get(cit, []), key=key)), "; ".join(sorted(mcsf.get(cit, []))),
                        "Rev. 5.1.1", "1.1", "NIST OLIR 110 and 109 (final, posted 2024-03-20)", h53, hcsf])
    print(f"{len(m53)} citations mapped to SP 800-53; {len(mcsf)} to CSF 1.1; sha256 {h53[:12]} {hcsf[:12]}")


if __name__ == "__main__":
    main()
