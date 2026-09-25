#!/usr/bin/env python3
"""Rebuild the CSF 2.0 core and official NIST crosswalks in 00_universal/.

Source: NIST CSF 2.0 Reference Tool download (all informative references).
Run when NIST publishes new informative references or a new CSF/800-53 release.
Usage: python3 tools/refresh_csf_crosswalks.py [path/to/local.xlsx]
"""
import csv, datetime, io, pathlib, re, sys, urllib.request

import openpyxl

URL = "https://csrc.nist.gov/extensions/nudp/services/json/csf/download?olirids=all"
ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / "00_universal"

# Informative-reference prefixes to export, mapped to output file names.
EXPORTS = {
    "SP 800-53 Rev 5.2.0": "crosswalks/csf2_to_sp800-53r5.csv",
    "SP 800-171 Rev 3": "crosswalks/csf2_to_sp800-171r3.csv",
}


def load_workbook(path=None):
    if path:
        return openpyxl.load_workbook(path, read_only=True)
    with urllib.request.urlopen(URL, timeout=120) as r:
        return openpyxl.load_workbook(io.BytesIO(r.read()), read_only=True)


def main():
    wb = load_workbook(sys.argv[1] if len(sys.argv) > 1 else None)
    rows = list(wb["CSF 2.0"].iter_rows(values_only=True))[2:]
    core, maps = [], {k: [] for k in EXPORTS}
    function = category = None
    for fn, cat, sub, _examples, refs in rows:
        if fn:
            function = re.match(r"(\w+) \((\w+)\)", fn).groups()
        if cat:
            m = re.match(r"(.+?) \((\w+\.\w+)\): (.*)", cat, re.S)
            category = (m.group(2), m.group(1).strip())
        if not sub:
            continue
        sid, text = sub.split(":", 1)
        sid = sid.strip()
        if text.strip().startswith("[Withdrawn"):
            continue  # CSF 1.1 items retained by NIST only for traceability
        core.append([function[1], function[0].title(), category[0], category[1], sid, " ".join(text.split())])
        for line in (refs or "").splitlines():
            if ":" not in line:
                continue
            src, ref = (s.strip() for s in line.split(":", 1))
            if src in maps:
                maps[src].append([sid, category[0], ref])

    stamp = datetime.date.today().isoformat()
    (OUT / "frameworks").mkdir(parents=True, exist_ok=True)
    with open(OUT / "frameworks/csf2_core.csv", "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["function_id", "function", "category_id", "category", "subcategory_id", "subcategory"])
        w.writerows(core)
    for src, rel in EXPORTS.items():
        (OUT / rel).parent.mkdir(parents=True, exist_ok=True)
        with open(OUT / rel, "w", newline="") as f:
            w = csv.writer(f, lineterminator="\n")
            w.writerow(["csf2_subcategory", "csf2_category", "reference", "reference_source", "retrieved"])
            w.writerows(r + [src, stamp] for r in maps[src])
    cats = {r[2] for r in core}
    assert len(cats) == 22 and len(core) == 106, (len(cats), len(core))  # CSF 2.0 shape
    print(f"{len(core)} subcategories; " + ", ".join(f"{s}: {len(v)} mappings" for s, v in maps.items()))


if __name__ == "__main__":
    main()
